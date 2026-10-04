"""agent.py — 可重複調用的本地 agent 核心（Ollama + 本機 GPU）

用法：
    python agent.py "把 ./reports 裡的檔案列出來，告訴我哪個最新"

在其他程式裡調用：
    from agent import Agent, ToolBox
    box = ToolBox()

    @box.tool
    def my_tool(x: str) -> str: ...

    answer = Agent("helper", "你是...", box).run("任務描述")
"""
from __future__ import annotations

import logging
import sys
from pathlib import Path
from typing import Callable

import ollama

log = logging.getLogger("agent")

MAX_TOOL_OUTPUT = 8000  # 工具回傳太長會吃光 context，超過就截斷


class ToolBox:
    """工具註冊表。所有自動化能力都寫成函式，註冊到這裡。"""

    def __init__(self) -> None:
        self._fns: dict[str, Callable] = {}
        self._needs_confirm: set[str] = set()

    def tool(self, fn: Callable | None = None, *, confirm: bool = False):
        """裝飾器。confirm=True 代表會寫入/寄出/刪除，執行前要人工同意。"""

        def register(f: Callable) -> Callable:
            self._fns[f.__name__] = f
            if confirm:
                self._needs_confirm.add(f.__name__)
            return f

        return register(fn) if fn else register

    def subset(self, *names: str) -> "ToolBox":
        """只挑幾個工具給某個 agent 用（工具越少，模型越不容易選錯）。"""
        box = ToolBox()
        for n in names:
            box._fns[n] = self._fns[n]
            if n in self._needs_confirm:
                box._needs_confirm.add(n)
        return box

    @property
    def functions(self) -> list[Callable]:
        return list(self._fns.values())

    def get(self, name: str) -> Callable | None:
        return self._fns.get(name)

    def needs_confirm(self, name: str) -> bool:
        return name in self._needs_confirm


def ask_in_terminal(name: str, args: dict) -> bool:
    """預設的確認方式：在終端機問 y/n。排程執行時換成 lambda *_: False。"""
    return input(f"\n[確認] 執行 {name}({args})？ (y/N) ").strip().lower() == "y"


class Agent:
    def __init__(
        self,
        name: str,
        system: str,
        toolbox: ToolBox,
        *,
        model: str = "gpt-oss:20b",
        num_ctx: int = 32768,
        max_steps: int = 12,
        host: str | None = None,
        approve: Callable[[str, dict], bool] = ask_in_terminal,
        client=None,
    ) -> None:
        self.name = name
        self.system = system
        self.toolbox = toolbox
        self.model = model
        self.num_ctx = num_ctx
        self.max_steps = max_steps
        self.approve = approve
        self.client = client or ollama.Client(host=host)

    def run(self, task: str) -> str:
        """丟一個任務進來，跑到模型給出最終答案為止。每次呼叫互相獨立。"""
        messages: list = [
            {"role": "system", "content": self.system},
            {"role": "user", "content": task},
        ]
        log.info("[%s] model=%s tools=%s", self.name, self.model,
                 [f.__name__ for f in self.toolbox.functions])
        for step in range(1, self.max_steps + 1):
            resp = self.client.chat(
                model=self.model,
                messages=messages,
                tools=self.toolbox.functions,
                options={"num_ctx": self.num_ctx},
            )
            msg = resp.message
            messages.append(msg)

            if not msg.tool_calls:
                return msg.content or ""

            for call in msg.tool_calls:
                name, args = call.function.name, dict(call.function.arguments or {})
                log.info("[%s] step %d → %s(%s)", self.name, step, name, args)
                result = self._execute(name, args)
                messages.append({"role": "tool", "tool_name": name, "content": result})

        return f"[{self.name}] 超過 {self.max_steps} 步仍未完成，已中止。"

    def _execute(self, name: str, args: dict) -> str:
        fn = self.toolbox.get(name)
        if fn is None:
            return f"ERROR: no tool named '{name}'."
        if self.toolbox.needs_confirm(name) and not self.approve(name, args):
            return "DENIED: the user did not approve this action."
        try:
            out = str(fn(**args))
        except Exception as e:  # 把錯誤回給模型，讓它自己修正參數重試
            log.warning("[%s] %s failed: %s", self.name, name, e)
            return f"ERROR: {type(e).__name__}: {e}"
        if len(out) > MAX_TOOL_OUTPUT:
            out = out[:MAX_TOOL_OUTPUT] + "\n...[truncated]"
        return out


# ---------------------------------------------------------------------------
# 範例工具：之後換成你自己的（查 SQL、讀 Excel、呼叫內部 API ...）
# ---------------------------------------------------------------------------
box = ToolBox()


@box.tool
def list_files(folder: str) -> str:
    """List files in a folder with size and last-modified time.

    Args:
        folder: Path of the folder to list.
    """
    rows = []
    for p in sorted(Path(folder).iterdir()):
        st = p.stat()
        rows.append(f"{p.name}\t{st.st_size} bytes\tmtime={int(st.st_mtime)}")
    return "\n".join(rows) or "(empty)"


@box.tool
def read_text_file(path: str) -> str:
    """Read a UTF-8 text file and return its content.

    Args:
        path: Path of the file to read.
    """
    return Path(path).read_text(encoding="utf-8")


TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".log", ".py", ".sql", ".html", ".xml", ".yaml", ".yml"}


@box.tool(confirm=True)
def write_text_file(path: str, content: str) -> str:
    """Write a PLAIN TEXT file (.txt .md .csv .json ...). Cannot create Excel files.

    Args:
        path: Path of the text file to write.
        content: Full text content to write.
    """
    ext = Path(path).suffix.lower()
    if ext in {".xlsx", ".xls", ".xlsm"}:
        raise ValueError("Excel files must be created with write_excel_sheet, not write_text_file.")
    if ext not in TEXT_EXTENSIONS:
        raise ValueError(f"'{ext}' is not a plain-text format. Allowed: {sorted(TEXT_EXTENSIONS)}")
    Path(path).write_text(content, encoding="utf-8")
    return f"wrote {len(content)} chars to {path}"


@box.tool(confirm=True)
def write_excel_sheet(path: str, sheet_name: str, headers: list[str], rows: list[list[str]]) -> str:
    """Create or replace ONE sheet in a real Excel .xlsx file. Call once per table/sheet.

    Args:
        path: Path of the .xlsx file, e.g. "output.xlsx".
        sheet_name: Name of the sheet to write, e.g. "Sales".
        headers: Column names, e.g. ["Product", "Qty"].
        rows: Table data, one inner list per row, e.g. [["Apple", 10], ["Banana", 5]].
    """
    from openpyxl import Workbook, load_workbook

    if Path(path).suffix.lower() != ".xlsx":
        raise ValueError("path must end with .xlsx")
    note = ""
    try:
        wb = load_workbook(path) if Path(path).exists() else None
    except Exception:  # 檔案存在但不是真正的 Excel（例如被寫成純文字），直接重建
        wb, note = None, " (existing file was not a valid Excel file and was replaced)"
    if wb is None:
        wb = Workbook()
        wb.remove(wb.active)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return f"wrote sheet '{sheet_name}' ({len(rows)} rows) to {path}{note}"


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    if len(sys.argv) < 2:
        sys.exit('用法: python agent.py "任務描述"')
    file_agent = Agent(
        name="file-helper",
        system=(
            "你是檔案助理。用提供的工具完成任務，最後用繁體中文簡短回報結果。"
            "Excel (.xlsx) 一律用 write_excel_sheet，每個工作表呼叫一次；"
            "write_text_file 只能寫純文字檔。"
        ),
        toolbox=box,
    )
    print(file_agent.run(" ".join(sys.argv[1:])))